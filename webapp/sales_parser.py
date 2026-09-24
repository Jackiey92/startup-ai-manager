"""Sales file parser: Excel ledgers and PDF contracts/orders -> structured data."""
from __future__ import annotations

import io
from datetime import date, datetime
from dataclasses import dataclass, field, asdict
from typing import Any

import openpyxl
import pymupdf


@dataclass
class StructuredTable:
    sheet: str
    headers: list[str]
    rows: list[dict[str, Any]]


@dataclass
class StructuredResult:
    format: str
    tables: list[StructuredTable] = field(default_factory=list)
    pdf_pages: list[dict[str, Any]] = field(default_factory=list)

    def to_dict(self) -> dict:
        d = {
            "format": self.format,
            "tables": [
                {"sheet": t.sheet, "headers": t.headers, "rows": t.rows}
                for t in self.tables
            ],
            "pdf_pages": self.pdf_pages,
        }
        return d


def parse_excel(blob: bytes) -> StructuredResult:
    wb = openpyxl.load_workbook(io.BytesIO(blob), data_only=True)
    result = StructuredResult(format="excel")

    for ws in wb.worksheets:
        raw_rows = [list(r) for r in ws.iter_rows(values_only=True)]
        non_empty = [r for r in raw_rows if any(v is not None and str(v).strip() for v in r)]
        if not non_empty:
            continue
        headers = [
            str(v).strip() if v is not None else f"列{i+1}"
            for i, v in enumerate(non_empty[0])
        ]
        rows = []
        for raw in non_empty[1:]:
            row = {}
            for i, value in enumerate(raw):
                key = headers[i] if i < len(headers) else f"列{i+1}"
                if value is not None and str(value).strip():
                    if isinstance(value, (date, datetime)):
                        value = value.isoformat()
                    row[key] = value
            if row:
                rows.append(row)
        result.tables.append(StructuredTable(sheet=ws.title, headers=headers, rows=rows))
    return result


def parse_pdf(blob: bytes) -> StructuredResult:
    result = StructuredResult(format="pdf")
    with pymupdf.open(stream=blob, filetype="pdf") as doc:
        for page_number, page in enumerate(doc, start=1):
            result.pdf_pages.append(
                {
                    "page": page_number,
                    "text": page.get_text("text", sort=True),
                }
            )
    return result


def parse_bytes(blob: bytes, filename: str) -> StructuredResult:
    name = filename.lower()
    if name.endswith((".xlsx", ".xlsm", ".xls")):
        return parse_excel(blob)
    if name.endswith(".pdf"):
        return parse_pdf(blob)
    raise ValueError(f"unsupported file: {filename}")

