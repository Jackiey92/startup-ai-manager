"""Build local synthetic Office fixtures containing one embedded picture each.

The source picture is supplied by the caller; generated Office files are test
artifacts only and are never uploaded or committed.
"""
from __future__ import annotations

import argparse
from pathlib import Path

from docx import Document
from docx.shared import Inches
from openpyxl import Workbook
from openpyxl.drawing.image import Image as XlsxImage
from pptx import Presentation
from pptx.util import Inches as PptxInches


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("image", type=Path)
    parser.add_argument("output", type=Path)
    args = parser.parse_args()
    args.output.mkdir(parents=True, exist_ok=True)

    doc = Document()
    doc.add_paragraph("合成 Office 图片验收：截图/印章占位图。")
    doc.add_picture(str(args.image), width=Inches(3))
    doc.save(args.output / "embedded-picture.docx")

    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    slide.shapes.add_picture(str(args.image), PptxInches(1), PptxInches(1), width=PptxInches(4))
    presentation.save(args.output / "embedded-picture.pptx")

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Evidence"
    sheet["A1"] = "合成 Office 图片验收"
    sheet.add_image(XlsxImage(str(args.image)), "B2")
    workbook.save(args.output / "embedded-picture.xlsx")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
