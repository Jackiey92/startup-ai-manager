from __future__ import annotations

import importlib.machinery
import importlib.util
from pathlib import Path

from docx import Document
from docx.shared import Inches
from openpyxl import Workbook
from openpyxl.drawing.image import Image as XlsxImage
from pptx import Presentation
from pptx.util import Inches as PptxInches


ROOT = Path(__file__).resolve().parents[1]
BRIDGE = ROOT / "skills" / "document-ingest" / "scripts" / "bridge"


def _bridge_module():
    loader = importlib.machinery.SourceFileLoader("sam_document_ingest_bridge", str(BRIDGE))
    spec = importlib.util.spec_from_loader(loader.name, loader)
    assert spec and spec.loader
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


def _make_png(path: Path) -> None:
    # A tiny valid RGB PNG; the image bytes are local test data only.
    from PIL import Image

    Image.new("RGB", (8, 8), (180, 20, 20)).save(path, format="PNG")


def _make_office(path: Path, image: Path, fmt: str) -> None:
    if fmt == "docx":
        document = Document()
        document.add_paragraph("No-suffix DOCX content")
        document.add_picture(str(image), width=Inches(1))
        document.save(path)
    elif fmt == "pptx":
        presentation = Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        slide.shapes.add_picture(str(image), PptxInches(1), PptxInches(1), width=PptxInches(1))
        presentation.save(path)
    elif fmt == "xlsx":
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Evidence"
        sheet["A1"] = "No-suffix XLSX content"
        sheet.add_image(XlsxImage(str(image)), "B2")
        workbook.save(path)
    else:  # pragma: no cover - the parametrization is deliberately narrow
        raise AssertionError(fmt)


def test_embedded_images_work_from_hash_only_office_objects(tmp_path: Path) -> None:
    """Content-addressed objects have no suffix, but Office images still extract."""
    bridge = _bridge_module()
    image = tmp_path / "fixture.png"
    _make_png(image)
    for fmt in ("docx", "pptx", "xlsx"):
        named = tmp_path / f"fixture.{fmt}"
        _make_office(named, image, fmt)
        hash_only = tmp_path / (f"{fmt}-" + "a" * 64)
        hash_only.write_bytes(named.read_bytes())
        images, warnings = bridge._extract_embedded_images(hash_only, "a" * 64, fmt)
        assert not warnings, (fmt, warnings)
        assert len(images) == 1, (fmt, images)
        assert images[0]["image_hash"]
        assert images[0]["source_loc"]["file_hash"] == "a" * 64


def test_image_extraction_failure_is_non_fatal(tmp_path: Path) -> None:
    bridge = _bridge_module()
    source = tmp_path / ("b" * 64)
    source.write_bytes(b"not an OOXML container")
    images, warnings = bridge._extract_embedded_images(source, "b" * 64, "xlsx")
    assert images == []
    assert warnings and "BadZipFile" in warnings[0]


def test_auto_office_prefers_docling_and_falls_back(monkeypatch, tmp_path: Path) -> None:
    bridge = _bridge_module()
    source = tmp_path / ("c" * 64)
    source.write_bytes(b"office bytes")
    calls: list[str] = []

    monkeypatch.setenv("SAM_INGEST_ENGINE", "auto")
    monkeypatch.setattr(bridge, "_engine_available", lambda name: True)
    monkeypatch.setattr(bridge, "_extract_embedded_images", lambda *args: ([], []))

    def fail_docling(*args):
        calls.append("docling")
        raise RuntimeError("simulated local conversion failure")

    def parse_mineru(*args):
        calls.append("mineru")
        return [], {"status": "parsed", "engine": "mineru", "warnings": []}

    monkeypatch.setattr(bridge, "_run_docling", fail_docling)
    monkeypatch.setattr(bridge, "_run_mineru", parse_mineru)
    result = bridge._manifest({"format": "xlsx", "filename": "report.xlsx"}, source)
    assert result["parse_summary"]["status"] == "parsed"
    assert result["parse_summary"]["engine"] == "mineru"
    assert calls == ["docling", "mineru"]
    assert "docling failed" in result["parse_summary"]["warnings"][0]


def test_mineru_status_treats_zero_exit_not_running_as_not_ready(monkeypatch) -> None:
    bridge = _bridge_module()

    monkeypatch.setattr(
        bridge.subprocess, "run",
        lambda command, **_kwargs: bridge.subprocess.CompletedProcess(
            command, 0, stdout="Server is not running.\n", stderr="",
        ),
    )
    ready, detail = bridge._mineru_status("mineru", {})
    assert ready is False
    assert "not running" in detail.lower()


def test_mineru_server_is_started_once_and_waited_until_ready(monkeypatch) -> None:
    bridge = _bridge_module()
    status_codes = iter((1, 1, 0))
    status_calls: list[list[str]] = []
    starts: list[dict] = []

    def fake_run(command, **kwargs):
        status_calls.append(command)
        return bridge.subprocess.CompletedProcess(command, next(status_codes), stdout="not ready", stderr="")

    class Starter:
        returncode = 0

        def communicate(self, timeout):
            assert timeout == 5
            return "started", ""

    def fake_popen(command, **kwargs):
        starts.append(kwargs)
        assert command == ["mineru", "server", "start"]
        return Starter()

    monkeypatch.setattr(bridge.subprocess, "run", fake_run)
    monkeypatch.setattr(bridge.subprocess, "Popen", fake_popen)
    monkeypatch.setattr(bridge.time, "sleep", lambda _seconds: None)
    bridge._ensure_mineru_server(mineru_bin="mineru", ready_timeout=5)
    assert status_calls == [["mineru", "server", "status"]] * 3
    assert len(starts) == 1 and starts[0]["start_new_session"] is True


def test_mineru_start_failure_is_reported_as_a_user_message(monkeypatch, tmp_path: Path) -> None:
    bridge = _bridge_module()
    source = tmp_path / "failure.pdf"
    source.write_bytes(b"not a PDF")
    monkeypatch.setattr(bridge, "_select_engines", lambda _fmt: (["mineru"], []))
    monkeypatch.setattr(bridge, "_extract_embedded_images", lambda *args: ([], []))

    def unavailable(*_args):
        raise bridge.MinerUServerUnavailable("server did not become ready")

    monkeypatch.setattr(bridge, "_run_mineru", unavailable)
    result = bridge._manifest({"format": "pdf", "filename": "failure.pdf"}, source)
    summary = result["parse_summary"]
    assert summary["status"] == "parse_failed"
    assert summary["user_message"] == "本地解析服务启动失败，请稍后重试或联系管理员。"
    assert "MinerUServerUnavailable" in summary["warnings"][0]


def test_mineru_parse_retries_until_the_basic_tier_is_ready(monkeypatch, tmp_path: Path) -> None:
    """A control-ready daemon may still be loading the parser model."""
    bridge = _bridge_module()
    source = tmp_path / "report.pdf"
    source.write_bytes(b"%PDF-1.4 local fixture")
    cache = tmp_path / "parsed.json"
    cache.write_text('{"pages": []}', encoding="utf-8")
    attempts: list[list[str]] = []

    monkeypatch.setattr(bridge, "_ensure_mineru_server", lambda **_kwargs: None)
    monkeypatch.setattr(bridge, "_mineru_cache_json", lambda *_args: cache)
    monkeypatch.setattr(bridge.time, "sleep", lambda _seconds: None)
    monkeypatch.setenv("SAM_MINERU_PARSE_READY_TIMEOUT", "5")

    def fake_run(command, **_kwargs):
        attempts.append(command)
        if len(attempts) == 1:
            return bridge.subprocess.CompletedProcess(
                command, 1, stdout="", stderr="quality_tier_unavailable",
            )
        return bridge.subprocess.CompletedProcess(command, 0, stdout='{"parse": {"tier": "basic"}}', stderr="")

    monkeypatch.setattr(bridge.subprocess, "run", fake_run)
    pages, summary = bridge._run_mineru(source, "d" * 64, "pdf", "report.pdf")
    assert pages == []
    assert summary["status"] == "parsed"
    assert len(attempts) == 2


def test_docling_provenance_keeps_the_engine_namespace() -> None:
    bridge = _bridge_module()

    class Item:
        self_ref = "#/texts/7"
        prov = []

    location = bridge._docling_provenance(Item(), "e" * 64, "docling:text/7")
    assert location["locator"] == "docling:text/7@#/texts/7"


def test_legacy_office_missing_converter_is_a_deployment_error(monkeypatch, tmp_path: Path) -> None:
    bridge = _bridge_module()
    source = tmp_path / "legacy.ppt"
    source.write_bytes(b"legacy binary")
    monkeypatch.setattr(bridge, "_soffice_binary", lambda: None)
    result = bridge._manifest({"format": "ppt", "filename": "legacy.ppt"}, source)
    summary = result["parse_summary"]
    assert summary["status"] == "engine_unavailable"
    assert "需安装 LibreOffice" in summary["user_message"]


def test_legacy_office_conversion_passes_ooxml_to_docling(monkeypatch, tmp_path: Path) -> None:
    bridge = _bridge_module()
    source = tmp_path / "legacy.doc"
    source.write_bytes(b"legacy binary")
    converter = tmp_path / "soffice"
    converter.write_text("placeholder", encoding="utf-8")
    converter.chmod(0o755)
    seen: dict[str, object] = {}

    def fake_run(command, **_kwargs):
        outdir = Path(command[command.index("--outdir") + 1])
        (outdir / "source.docx").write_bytes(b"ooxml")
        return bridge.subprocess.CompletedProcess(command, 0, stdout="converted", stderr="")

    def fake_docling(path, _file_hash, fmt, filename):
        seen.update({"path": path, "fmt": fmt, "filename": filename})
        return [], {"status": "parsed", "engine": "docling", "warnings": []}

    monkeypatch.setattr(bridge, "_soffice_binary", lambda: str(converter))
    monkeypatch.setattr(bridge.subprocess, "run", fake_run)
    monkeypatch.setattr(bridge, "_select_engines", lambda fmt: (["docling"], []) if fmt == "docx" else ([], []))
    monkeypatch.setattr(bridge, "_extract_embedded_images", lambda *_args: ([], []))
    monkeypatch.setattr(bridge, "_run_docling", fake_docling)
    result = bridge._manifest({"format": "doc", "filename": "legacy.doc"}, source)
    assert result["parse_summary"]["status"] == "parsed"
    assert seen["fmt"] == "docx"
    assert seen["filename"] == "legacy.docx"
    assert any("converted legacy .doc" in item for item in result["parse_summary"]["warnings"])
